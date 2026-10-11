"""Carga del padron SSCO (Sujetos Sin Capacidad Operativa) como marca 'ssco'.

Uso:  python cargar_ssco.py RUTA.xlsx --corte 2026-09-30 --dry-run

Formato real (corte 30-09-2026, 766 filas):
- El RUC viene como NUMERO decimal de Excel (10002179496.0): se pasa a texto de 11 digitos.
- Documento del representante: numero de 11 digitos (RUC), o texto con prefijo en varias
  grafias ("DNI: 8d", "DNI:8d", "DNI N.º 8d", "CE: 9d", "CE N.º 9d", "CE N° 9d"; tambien
  PAS / CED). Un DNI puede venir sin su cero inicial (1538115 -> 01538115).
- Varios representantes: documentos y nombres separados por comas, emparejados por posicion.
- El documento se guarda y se muestra TAL CUAL lo publica SUNAT (decision de Duilio).

--dry-run no escribe nada: reporta y guarda cargas-padron/dryrun_ssco_<corte>.json.
"""
from __future__ import annotations

import argparse
import json
import os
import re
import sys
import time
from collections import Counter
from datetime import date, datetime

from app import padron as P

sys.stdout.reconfigure(encoding="utf-8")
CARPETA = os.path.join(os.path.dirname(os.path.abspath(__file__)), "cargas-padron")
_PREFIJO = re.compile(r"^(DNI|CE|PAS|CED|RUC)\s*(?:N\s*[.°º]*\s*[°º]?\.?)?\s*:?\s*([A-Z0-9-]+)$", re.I)


def texto_numero(v) -> str:
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, int):
        return str(v)
    return str(v or "").strip()


def parsear_documento(parte: str) -> dict:
    """'DNI N.º 76969380' -> {tipo:'DNI', doc:'76969380', publicado:'DNI N.º 76969380'}."""
    publicado = parte.strip()
    m = _PREFIJO.match(publicado)
    if m:
        tipo, doc = m.group(1).upper(), m.group(2)
    elif publicado.isdigit():
        tipo, doc = ("RUC", publicado) if len(publicado) == 11 else ("DNI", publicado) if len(publicado) <= 8 else ("?", publicado)
    else:
        tipo, doc = "?", publicado
    if tipo == "DNI" and doc.isdigit() and len(doc) < 8:
        doc = doc.zfill(8)                     # Excel le comio el cero inicial
    return {"tipo": tipo, "doc": doc, "publicado": publicado}


def representantes(doc_celda, nombres_celda) -> tuple[list[dict], bool]:
    """Empareja documentos y nombres por posicion. -> (lista, cuadra)"""
    docs = [x for x in (texto_numero(doc_celda) if not isinstance(doc_celda, str) else doc_celda).split(",") if x.strip()]
    noms = [x.strip() for x in str(nombres_celda or "").split(",") if x.strip()]
    cuadra = len(docs) == len(noms)
    out = []
    for i, d in enumerate(docs):
        r = parsear_documento(d)
        r["nombre"] = noms[i] if cuadra else (str(nombres_celda or "").strip() if i == 0 else "")
        out.append(r)
    return out, cuadra


def fecha(v):
    return v.date().isoformat() if isinstance(v, datetime) else (v.isoformat() if isinstance(v, date) else None)


def leer(ruta):
    from openpyxl import load_workbook
    ws = load_workbook(ruta, read_only=True, data_only=True).worksheets[0]
    filas = [f for f in ws.iter_rows(values_only=True) if any(v is not None for v in f)]
    cab = [str(v or "").strip() for v in filas[0]]
    if not cab[0].upper().startswith("RUC") or len(cab) < 9:
        raise SystemExit(f"Cabecera inesperada: {cab}")
    return cab, filas[1:]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx"); ap.add_argument("--corte", required=True, type=date.fromisoformat)
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument("--dry-run", action="store_true"); g.add_argument("--prueba", action="store_true")
    g.add_argument("--aplicar", action="store_true")
    a = ap.parse_args()
    cab, filas = leer(a.xlsx)
    marcas, problemas, tipos, n_multi = [], [], Counter(), 0
    vistos = set()
    for i, f in enumerate(filas, 2):
        ruc = texto_numero(f[0])
        if not P.RUC_RE.match(ruc):
            problemas.append({"fila": i, "problema": "ruc_invalido", "valor": ruc}); continue
        if ruc in vistos:
            problemas.append({"fila": i, "problema": "duplicado", "valor": ruc}); continue
        vistos.add(ruc)
        reps, cuadra = representantes(f[6], f[7])
        if not cuadra:
            problemas.append({"fila": i, "problema": "docs_y_nombres_no_cuadran", "valor": f"{f[6]} / {f[7]}"})
        if len(reps) > 1:
            n_multi += 1
        for r in reps:
            tipos[r["tipo"]] += 1
            if r["tipo"] == "?":
                problemas.append({"fila": i, "problema": "documento_sin_tipo", "valor": r["publicado"]})
        if not P.ruc_dv_ok(ruc):
            problemas.append({"fila": i, "problema": "dv_no_cuadra(se_carga)", "valor": ruc})
        marcas.append({"ruc": ruc, "padron": "ssco", "desde": fecha(f[5]), "corte": a.corte.isoformat(),
                       "detalle": {"razon_social": str(f[1] or "").strip(), "domicilio": str(f[2] or "").strip(),
                                   "resolucion": str(f[3] or "").strip(), "emision": fecha(f[4]), "firme": fecha(f[5]),
                                   "publicacion": fecha(f[8]), "representantes": reps}})
    rep = {"modo": "dry-run", "archivo": a.xlsx, "corte": a.corte.isoformat(), "generado": time.strftime("%Y-%m-%d %H:%M:%S"),
           "cabecera": cab, "filas": len(filas), "marcas": len(marcas), "con_varios_representantes": n_multi,
           "tipos_de_documento": dict(tipos.most_common()), "problemas": problemas,
           "publicacion_por_mes": dict(sorted(Counter((m["detalle"]["publicacion"] or "")[:7] for m in marcas).items())),
           "ejemplos": [m for m in marcas if len(m["detalle"]["representantes"]) > 1][:2] + marcas[:2]
                       + [m for m in marcas if any(r["publicado"] != r["doc"] for r in m["detalle"]["representantes"])][:4]}
    if a.prueba or a.aplicar:
        import asyncio
        asyncio.run(aplicar("prueba_padron" if a.prueba else "public", marcas, a.corte))
        return
    os.makedirs(CARPETA, exist_ok=True)
    salida = os.path.join(CARPETA, f"dryrun_ssco_{a.corte.isoformat()}.json")
    json.dump(rep, open(salida, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(f"{len(filas)} filas → {len(marcas)} marcas SSCO; {n_multi} con varios representantes; "
          f"documentos {dict(tipos.most_common())}; {len(problemas)} observaciones. Reporte: {salida}")


def url_bd() -> str:
    for l in open(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".env"), encoding="utf-8"):
        m = re.match(r"^DATABASE_URL=(.+)", l.strip())
        if m:
            return m[1].strip().strip('"')
    raise SystemExit("Falta DATABASE_URL en .env")


async def aplicar(esquema: str, marcas: list[dict], corte: date):
    """Reemplaza las marcas 'ssco' por las del corte y deja en el historial quien entra,
    sale o cambia. Todo en una transaccion."""
    import asyncpg
    sql = open(os.path.join(os.path.dirname(os.path.abspath(__file__)), "sql", "alter_v12_verificador.sql"),
               encoding="utf-8").read()
    conn = await asyncpg.connect(url_bd())
    try:
        if esquema != "public":
            await conn.execute(f"CREATE SCHEMA IF NOT EXISTS {esquema}")
        await conn.execute(f"SET search_path TO {esquema}, public")
        await conn.execute(sql)
        previas = {r["ruc"]: json.loads(r["detalle"]) for r in
                   await conn.fetch("SELECT ruc, detalle FROM ruc_marcas WHERE padron = 'ssco'")}
        nuevas = {m["ruc"]: m for m in marcas}
        entra = [r for r in nuevas if r not in previas]
        sale = [r for r in previas if r not in nuevas]
        cambia = [r for r in nuevas if r in previas and previas[r] != nuevas[r]["detalle"]]
        async with conn.transaction():
            hist = ([(r, "entra", nuevas[r]["detalle"]) for r in entra] + [(r, "sale", previas[r]) for r in sale]
                    + [(r, "cambia", nuevas[r]["detalle"]) for r in cambia])
            await conn.executemany(
                "INSERT INTO ruc_marcas_hist (ruc, padron, evento, corte, detalle) VALUES ($1, 'ssco', $2, $3, $4::jsonb)",
                [(r, ev, corte, json.dumps(d, ensure_ascii=False)) for r, ev, d in hist])
            await conn.execute("DELETE FROM ruc_marcas WHERE padron = 'ssco'")
            await conn.executemany(
                "INSERT INTO ruc_marcas (ruc, padron, desde, corte, detalle) VALUES ($1, 'ssco', $2, $3, $4::jsonb)",
                [(m["ruc"], date.fromisoformat(m["desde"]) if m["desde"] else None, corte,
                  json.dumps(m["detalle"], ensure_ascii=False)) for m in marcas])
            await conn.execute("UPDATE padrones SET corte = $1 WHERE codigo = 'ssco'", corte)
        print(f"[{esquema}] SSCO al {corte}: {len(marcas)} marcas vigentes · historial: "
              f"{len(entra)} entran, {len(sale)} salen, {len(cambia)} cambian.")
    finally:
        await conn.close()


if __name__ == "__main__":
    main()
