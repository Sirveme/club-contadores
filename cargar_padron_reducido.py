"""Carga del Padron Reducido de SUNAT (listado general + locales anexos).

Uso (lo corre Duilio al descargar cada corte):
  python cargar_padron_reducido.py --ruc RUTA.txt --anexos RUTA.txt --corte 2026-10-10 \\
         --corte-anexos 2026-10-10 [--cruzar-ssco SSCO.xlsx] --dry-run

Modos:
  --dry-run   Lee los archivos completos por streaming y reporta conteos, rechazos y
              ejemplos. NO escribe en la base (solo lee cat_ubigeo). Guarda el reporte y
              la huella de los archivos en cargas-padron/ (ignorado por git).
  --prueba    (siguiente paso, tras aprobar el dry-run) carga en el esquema prueba_padron.
  --aplicar   (despues) carga en produccion: tabla nueva + renombrado atomico + bak_.
              Exigira un --dry-run previo con la MISMA huella de archivos.
"""
from __future__ import annotations

import argparse
import asyncio
import hashlib
import json
import os
import re
import sys
import time
from collections import Counter
from datetime import date

from app import padron as P

sys.stdout.reconfigure(encoding="utf-8")
CARPETA = os.path.join(os.path.dirname(os.path.abspath(__file__)), "cargas-padron")


def url_bd() -> str:
    for l in open(os.path.join(os.path.dirname(os.path.abspath(__file__)), ".env"), encoding="utf-8"):
        m = re.match(r"^DATABASE_URL=(.+)", l.strip())
        if m:
            return m[1].strip().strip('"')
    raise SystemExit("Falta DATABASE_URL en .env")


async def cargar_ubigeos() -> dict:
    import asyncpg
    c = await asyncpg.connect(url_bd())
    rows = await c.fetch("SELECT ubigeo, departamento, provincia, distrito FROM cat_ubigeo")
    await c.close()
    return {r["ubigeo"]: (r["departamento"], r["provincia"], r["distrito"]) for r in rows}


class Bitmap:
    """Presencia de RUC sin guardar 18 M de strings: un bit por los 9 digitos tras el
    prefijo (125 MB por prefijo). Incluye el verificador: hay RUC que solo difieren en el."""
    def __init__(self):
        self.m = {p: bytearray(125_000_000) for p in ("10", "15", "17", "20")}

    def _pos(self, ruc):
        n = int(ruc[2:]); return self.m[ruc[:2]], n >> 3, 1 << (n & 7)

    def agregar(self, ruc) -> bool:
        """True si ya estaba (duplicado)."""
        b, i, bit = self._pos(ruc)
        ya = b[i] & bit
        b[i] |= bit
        return bool(ya)

    def tiene(self, ruc) -> bool:
        b, i, bit = self._pos(ruc); return bool(b[i] & bit)


def leer_ssco(ruta) -> set[str]:
    from openpyxl import load_workbook
    ws = load_workbook(ruta, read_only=True, data_only=True).worksheets[0]
    out = set()
    for f in ws.iter_rows(min_row=2, values_only=True):
        if f and f[0] is not None:
            v = f[0]
            out.add(str(int(v)) if isinstance(v, float) else str(v).strip())
    return out


def ejemplos(lista, n, x):
    if len(lista) < n:
        lista.append(x)


def dry_run_listado(ruta, ubis, ssco: set[str]):
    t0 = time.time()
    sha = hashlib.sha256()
    vistos = Bitmap()
    r = Counter(); estados = Counter(); conds = Counter(); semaf = Counter(); prefijos = Counter()
    sin_mapa_e = Counter(); sin_mapa_c = Counter()
    ej_rechazo, ej_escapado, ej_dom, ej_dv = {}, [], [], []
    ubigeo_desconocido = Counter()
    ssco_vistos = {}
    largo_rs = 0; bytes_dom = 0; bytes_rs = 0
    with open(ruta, "rb") as f:
        cab = f.readline(); sha.update(cab)
        for n, raw in enumerate(f, 1):
            sha.update(raw)
            lin = raw.decode("latin-1")
            c = P.partir(lin)
            if len(c) != 15:
                r["columnas"] += 1; ej_rechazo.setdefault("columnas", []).append(lin.strip()[:160]) if len(ej_rechazo.get("columnas", [])) < 5 else None
                continue
            ruc = c[0].strip()
            if not P.RUC_RE.match(ruc):
                r["ruc_invalido"] += 1
                lst = ej_rechazo.setdefault("ruc_invalido", []); ejemplos(lst, 5, f"{ruc} | {c[1][:50]}")
                continue
            if vistos.agregar(ruc):
                r["duplicado"] += 1
                lst = ej_rechazo.setdefault("duplicado", []); ejemplos(lst, 5, ruc)
                continue
            if "\\" in lin:
                ejemplos(ej_escapado, 8, {"ruc": ruc, "razon_social": c[1], "domicilio": P.domicilio_listado(c, ubis)})
            if not P.ruc_dv_ok(ruc):
                r["dv_no_cuadra(se_carga)"] += 1; ejemplos(ej_dv, 5, f"{ruc} | {c[1][:50]}")
            r["aceptadas"] += 1
            prefijos[ruc[:2]] += 1
            est, cond = c[2].strip(), P.vacio(c[3])
            estados[est] += 1; conds[cond] += 1
            if est not in P.ESTADOS: sin_mapa_e[est] += 1
            if cond not in P.CONDICIONES: sin_mapa_c[cond] += 1
            semaf[P.semaforo_fila(est, cond)] += 1
            c[1] = P.razon_social(c[1])
            if not c[1]:
                r["sin_razon_social(se_carga)"] += 1
            largo_rs = max(largo_rs, len(c[1])); bytes_rs += len(c[1])
            ubi = P.vacio(c[4])
            if ubi and ubi not in ubis: ubigeo_desconocido[ubi] += 1
            dom = P.domicilio_listado(c, ubis)
            if dom:
                r["con_domicilio"] += 1; bytes_dom += len(dom)
                if n % 104729 == 0: ejemplos(ej_dom, 20, {"ruc": ruc, "domicilio": dom})
            if ruc in ssco:
                ssco_vistos[ruc] = (est, cond)
            if n % 3_000_000 == 0:
                print(f"  ... {n:,} filas ({time.time() - t0:.0f}s)", flush=True)
    ssco_estado = Counter(f"{P.ESTADOS.get(e, (e,))[0]} · {P.CONDICIONES.get(c_, (c_,))[0]}" for e, c_ in ssco_vistos.values())
    est_bytes = r["aceptadas"] * (11 + 40 + 28) + bytes_rs + bytes_dom
    return {
        "archivo": ruta, "bytes": os.path.getsize(ruta), "sha256": sha.hexdigest(), "filas": n,
        "segundos": round(time.time() - t0),
        "resultado": dict(r), "prefijos": dict(prefijos.most_common()),
        "estados": {k: {"filas": v, "nombre": P.ESTADOS.get(k, ("¿?",))[0], "semaforo": P.ESTADOS.get(k, ("", "¿?"))[1]}
                    for k, v in estados.most_common()},
        "condiciones": {(k or "-"): {"filas": v, "nombre": P.CONDICIONES.get(k, ("¿?",))[0], "semaforo": P.CONDICIONES.get(k, ("", "¿?"))[1]}
                        for k, v in conds.most_common()},
        "SIN_MAPEO_estados": dict(sin_mapa_e), "SIN_MAPEO_condiciones": dict(sin_mapa_c),
        "semaforo_por_fila": dict(semaf.most_common()),
        "razon_social_max": largo_rs, "ubigeos_desconocidos": dict(ubigeo_desconocido.most_common(10)),
        "ejemplos_rechazo": ej_rechazo, "ejemplos_pipe_escapado": ej_escapado,
        "ejemplos_dv": ej_dv, "ejemplos_domicilio": ej_dom,
        "ssco": {"en_archivo_ssco": len(ssco), "encontrados_en_padron": len(ssco_vistos),
                 "no_encontrados": sorted(ssco - set(ssco_vistos))[:20], "estado_de_los_ssco": dict(ssco_estado.most_common())},
        "estimado_tabla_GB": round(est_bytes * 1.35 / 1e9, 2),
    }, vistos


def dry_run_anexos(ruta, ubis, vistos: Bitmap):
    t0 = time.time()
    sha = hashlib.sha256()
    r = Counter(); por_ruc = Counter(); ej_dom = []; huerfanos = []
    with open(ruta, "rb") as f:
        cab = f.readline(); sha.update(cab)
        for n, raw in enumerate(f, 1):
            sha.update(raw)
            c = P.partir(raw.decode("latin-1"))
            if len(c) != 12:
                r["columnas"] += 1; continue
            ruc = c[0].strip()
            if not P.RUC_RE.match(ruc):
                r["ruc_invalido"] += 1; continue
            dom = P.domicilio_anexo(c, ubis)
            if not dom:
                r["sin_domicilio"] += 1; continue
            if not vistos.tiene(ruc):
                r["ruc_no_esta_en_listado(se_carga)"] += 1; ejemplos(huerfanos, 5, ruc)
            r["aceptadas"] += 1; por_ruc[ruc] += 1
            if n % 37813 == 0: ejemplos(ej_dom, 12, {"ruc": ruc, "domicilio": dom})
    dist = Counter(min(v, 10) for v in por_ruc.values())
    return {"archivo": ruta, "bytes": os.path.getsize(ruta), "sha256": sha.hexdigest(), "filas": n,
            "segundos": round(time.time() - t0), "resultado": dict(r), "rucs_con_anexos": len(por_ruc),
            "anexos_por_ruc(10=10+)": dict(sorted(dist.items())), "max_anexos_un_ruc": max(por_ruc.values()),
            "ejemplos_domicilio": ej_dom, "ejemplos_ruc_no_en_listado": huerfanos}


SQL_V12 = os.path.join(os.path.dirname(os.path.abspath(__file__)), "sql", "alter_v12_verificador.sql")
LOTE = 50_000


async def _copiar(conn, esquema, tabla, columnas, filas_iter, cada=1_000_000):
    """COPY por lotes desde un iterador de tuplas. -> filas copiadas."""
    lote, total, t0 = [], 0, time.time()
    for fila in filas_iter:
        lote.append(fila)
        if len(lote) >= LOTE:
            await conn.copy_records_to_table(tabla, records=lote, columns=columnas, schema_name=esquema)
            total += len(lote); lote = []
            if total % cada < LOTE:
                print(f"  ... {total:,} filas subidas ({time.time() - t0:.0f}s)", flush=True)
    if lote:
        await conn.copy_records_to_table(tabla, records=lote, columns=columnas, schema_name=esquema)
        total += len(lote)
    return total


async def _intercambiar(conn, base: str, sufijo_bak: str, indices: dict[str, str]):
    """<base>_nuevo pasa a ser <base> y la anterior queda como bak_<sufijo>_<base>, con sus
    indices renombrados igual ({indice_nuevo: indice_final}). En una transaccion: o cambia
    completo o no cambia nada."""
    bak = f"bak_{sufijo_bak}_{base}"
    async with conn.transaction():
        existe = await conn.fetchval("SELECT to_regclass($1) IS NOT NULL", base)
        if existe:
            await conn.execute(f"DROP TABLE IF EXISTS {bak}")
            await conn.execute(f"ALTER TABLE {base} RENAME TO {bak}")
            for final in indices.values():
                if await conn.fetchval("SELECT to_regclass($1) IS NOT NULL", final):
                    await conn.execute(f"ALTER INDEX {final} RENAME TO bak_{sufijo_bak}_{final}")
        await conn.execute(f"ALTER TABLE {base}_nuevo RENAME TO {base}")
        for nuevo, final in indices.items():
            await conn.execute(f"ALTER INDEX {nuevo} RENAME TO {final}")
    return bak if existe else None


def _filas_listado(ruta, ubis, sha, incluir):
    with open(ruta, "rb") as f:
        sha.update(f.readline())
        for raw in f:
            sha.update(raw)
            c = P.partir(raw.decode("latin-1"))
            if len(c) != 15:
                continue
            ruc = c[0].strip()
            if not P.RUC_RE.match(ruc) or not incluir(ruc, c):
                continue
            ubi = P.vacio(c[4]) or None
            yield (ruc, P.razon_social(c[1]), c[2].strip(), P.vacio(c[3]), ubi, P.domicilio_listado(c, ubis))


def _filas_anexos(ruta, ubis, sha, incluir_ruc):
    with open(ruta, "rb") as f:
        sha.update(f.readline())
        for raw in f:
            sha.update(raw)
            c = P.partir(raw.decode("latin-1"))
            if len(c) != 12:
                continue
            ruc = c[0].strip()
            if not P.RUC_RE.match(ruc) or not incluir_ruc(ruc):
                continue
            dom = P.domicilio_anexo(c, ubis)
            if dom:
                yield (ruc, P.vacio(c[1]) or None, dom)


async def cargar(a, esquema: str):
    """--prueba (esquema prueba_padron, con muestra) y --aplicar (public, completo)."""
    import asyncpg
    prueba = esquema != "public"
    reporte = os.path.join(CARPETA, f"dryrun_padron_{a.corte.isoformat()}.json")
    dry = json.load(open(reporte, encoding="utf-8")) if os.path.exists(reporte) else None
    if not prueba and not dry:
        raise SystemExit(f"--aplicar exige el --dry-run previo de este corte ({reporte}).")
    conn = await asyncpg.connect(url_bd(), command_timeout=None)
    try:
        if prueba:
            await conn.execute(f"CREATE SCHEMA IF NOT EXISTS {esquema}")
        await conn.execute(f"SET search_path TO {esquema}, public")
        await conn.execute(open(SQL_V12, encoding="utf-8").read())
        ubis = {r["ubigeo"]: (r["departamento"], r["provincia"], r["distrito"])
                for r in await conn.fetch("SELECT ubigeo, departamento, provincia, distrito FROM public.cat_ubigeo")}
        ubis.update({k: v for k, v in P.UBIGEOS_COMPLEMENTO.items() if k not in ubis})
        incluidos = Bitmap() if prueba else None
        if prueba:
            # Muestra: todo Loreto, 1 de cada 20 RUC del resto, los SSCO y los negocios nuevos.
            extra = leer_ssco(a.cruzar_ssco) if a.cruzar_ssco else set()
            extra |= {r["ruc"] for r in await conn.fetch("SELECT ruc FROM public.nuevos_negocios")}

            def incluir(ruc, c):
                ok = ruc in extra or P.vacio(c[4]).startswith("16") or int(ruc[2:10]) % 20 == 0
                if ok:
                    incluidos.agregar(ruc)
                return ok
        else:
            def incluir(ruc, c):
                return True
        sufijo = time.strftime("%Y%m%d")

        print(f"[{esquema}] Listado general → padron_reducido")
        await conn.execute("DROP TABLE IF EXISTS padron_reducido_stg")
        await conn.execute("CREATE UNLOGGED TABLE padron_reducido_stg (ruc text, razon_social text, estado text, "
                           "condicion text, ubigeo text, domicilio text)")
        sha = hashlib.sha256()
        n = await _copiar(conn, esquema, "padron_reducido_stg",
                          ["ruc", "razon_social", "estado", "condicion", "ubigeo", "domicilio"],
                          _filas_listado(a.ruc, ubis, sha, incluir))
        if not prueba and sha.hexdigest() != dry["listado"]["sha256"]:
            await conn.execute("DROP TABLE padron_reducido_stg")
            raise SystemExit("El archivo NO es el mismo del --dry-run aprobado (huella distinta). No se cambió nada.")
        await conn.execute("DROP TABLE IF EXISTS padron_reducido_nuevo")
        await conn.execute(f"""
            CREATE TABLE padron_reducido_nuevo AS
            SELECT DISTINCT ON (ruc) ruc, razon_social, estado, condicion, ubigeo::char(6) AS ubigeo, domicilio,
                   DATE '{a.corte.isoformat()}' AS corte
            FROM padron_reducido_stg ORDER BY ruc""")
        await conn.execute("ALTER TABLE padron_reducido_nuevo ADD CONSTRAINT padron_reducido_nuevo_pkey PRIMARY KEY (ruc)")
        await conn.execute("DROP TABLE padron_reducido_stg")
        bak = await _intercambiar(conn, "padron_reducido", sufijo, {"padron_reducido_nuevo_pkey": "padron_reducido_pkey"})
        await conn.execute("ANALYZE padron_reducido")
        print(f"  {n:,} filas cargadas" + (f"; la anterior quedó como {bak}" if bak else ""))

        print(f"[{esquema}] Locales anexos → padron_anexos")
        await conn.execute("DROP TABLE IF EXISTS padron_anexos_stg")
        await conn.execute("CREATE UNLOGGED TABLE padron_anexos_stg (ruc text, ubigeo text, domicilio text)")
        sha_a = hashlib.sha256()
        na = await _copiar(conn, esquema, "padron_anexos_stg", ["ruc", "ubigeo", "domicilio"],
                           _filas_anexos(a.anexos, ubis, sha_a, (incluidos.tiene if prueba else (lambda r: True))))
        if not prueba and sha_a.hexdigest() != dry["anexos"]["sha256"]:
            await conn.execute("DROP TABLE padron_anexos_stg")
            raise SystemExit("El archivo de anexos NO es el mismo del --dry-run aprobado. El listado ya se cambió; "
                             "los anexos no. Revisa y vuelve a correr.")
        await conn.execute("DROP TABLE IF EXISTS padron_anexos_nuevo")
        await conn.execute(f"""
            CREATE TABLE padron_anexos_nuevo AS
            SELECT ruc, ubigeo::char(6) AS ubigeo, domicilio, DATE '{a.corte_anexos.isoformat()}' AS corte
            FROM padron_anexos_stg ORDER BY ruc""")
        await conn.execute("CREATE INDEX ix_anexos_ruc_nuevo ON padron_anexos_nuevo (ruc)")
        await conn.execute("DROP TABLE padron_anexos_stg")
        bak_a = await _intercambiar(conn, "padron_anexos", sufijo, {"ix_anexos_ruc_nuevo": "ix_anexos_ruc"})
        await conn.execute("ANALYZE padron_anexos")
        print(f"  {na:,} anexos cargados" + (f"; los anteriores quedaron como {bak_a}" if bak_a else ""))
        await conn.execute("UPDATE padrones SET corte = $1 WHERE codigo = 'padron_reducido'", a.corte)
        await conn.execute("UPDATE padrones SET corte = $1 WHERE codigo = 'padron_anexos'", a.corte_anexos)
        tam = await conn.fetchval("SELECT pg_size_pretty(pg_total_relation_size('padron_reducido') + "
                                  "pg_total_relation_size('padron_anexos'))")
        print(f"Listo en {esquema}: {n:,} RUC + {na:,} anexos ({tam}).")
    finally:
        await conn.close()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--ruc", required=True); ap.add_argument("--anexos", required=True)
    ap.add_argument("--corte", required=True, type=date.fromisoformat)
    ap.add_argument("--corte-anexos", required=True, type=date.fromisoformat)
    ap.add_argument("--cruzar-ssco")
    g = ap.add_mutually_exclusive_group(required=True)
    g.add_argument("--dry-run", action="store_true"); g.add_argument("--prueba", action="store_true")
    g.add_argument("--aplicar", action="store_true")
    a = ap.parse_args()
    if a.prueba or a.aplicar:
        asyncio.run(cargar(a, "prueba_padron" if a.prueba else "public"))
        return
    print("Leyendo catálogo de ubigeos (solo lectura)…")
    ubis = asyncio.run(cargar_ubigeos())
    ubis.update({k: v for k, v in P.UBIGEOS_COMPLEMENTO.items() if k not in ubis})
    ssco = leer_ssco(a.cruzar_ssco) if a.cruzar_ssco else set()
    print(f"Listado general: {a.ruc}")
    listado, vistos = dry_run_listado(a.ruc, ubis, ssco)
    print(f"Locales anexos: {a.anexos}")
    anexos = dry_run_anexos(a.anexos, ubis, vistos)
    rep = {"modo": "dry-run", "corte": a.corte.isoformat(), "corte_anexos": a.corte_anexos.isoformat(),
           "generado": time.strftime("%Y-%m-%d %H:%M:%S"), "listado": listado, "anexos": anexos}
    os.makedirs(CARPETA, exist_ok=True)
    salida = os.path.join(CARPETA, f"dryrun_padron_{a.corte.isoformat()}.json")
    json.dump(rep, open(salida, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    L = listado
    print(f"\nListo. {L['filas']:,} filas → {L['resultado'].get('aceptadas', 0):,} aceptadas en {L['segundos']} s; "
          f"anexos {anexos['resultado'].get('aceptadas', 0):,}. Reporte: {salida}")


if __name__ == "__main__":
    main()
